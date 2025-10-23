#!/usr/bin/env python3
"""
Enhanced AI Template Populator for SDTM data
Maps adae.xlsx columns to narrative template placeholders
"""

import os
import pandas as pd
import openai
from docx import Document
from docx.shared import Inches
import json
import re
from typing import Dict, List, Any
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class EnhancedAITemplatePopulator:
    def __init__(self, api_key: str = None):
        self.api_key = api_key or os.getenv('OPENAI_API_KEY')
        if not self.api_key:
            raise ValueError("OpenAI API key not provided. Set OPENAI_API_KEY environment variable")
        
        openai.api_key = self.api_key
        self.client = openai.OpenAI(api_key=self.api_key)
        
        # Mapping from adae.xlsx columns to template placeholders
        self.column_mapping = {
            'study_identifier': 'Study Identifier',
            'unique_subject_identifier': 'Unique Subject Identifier',
            'age': 'Age',
            'age_units': 'Age Units',
            'sex': 'Sex',
            'race': 'Race',
            'ethnicity': 'Ethnicity',
            'study_site_identifier': 'Study Site Identifier',
            'phase': 'Phase',
            'description_of_planned_arm': 'Description of Planned Arm',
            'description_of_actual_arm': 'Description of Actual Arm',
            'datetime_of_first_exposure_to_treatment': 'Datetime of First Exposure to Treatment',
            'datetime_of_last_exposure_to_treatment': 'Datetime of Last Exposure to Treatment',
            'reported_term_for_the_adverse_event': 'Reported Term for the Adverse Event',
            'dictionary_derived_term': 'Dictionary-Derived Term',
            'preferred_term_code': 'Preferred Term Code',
            'body_system_or_organ_class': 'Body System or Organ Class',
            'start_datetime_of_adverse_event': 'Start Date/Time of Adverse Event',
            'end_datetime_of_adverse_event': 'End Date/Time of Adverse Event',
            'analysis_start_relative_day': 'Analysis Start Relative Day',
            'analysis_end_relative_day': 'Analysis End Relative Day',
            'serious_event': 'Serious Event',
            'standard_toxicity_grade': 'Standard Toxicity Grade',
            'treatment_emergent_flag': 'Treatment Emergent Flag',
            'results_in_death': 'Results in Death',
            'is_life_threatening': 'Is Life Threatening',
            'requires_or_prolongs_hospitalization': 'Requires or Prolongs Hospitalization',
            'persist_or_signif_disabilityincapacity': 'Persist or Signif Disability/Incapacity',
            'other_medically_important_serious_event': 'Other Medically Important Serious Event',
            'action_taken_with_study_treatment': 'Action Taken with Study Treatment',
            'causality': 'Causality',
            'analysis_causality': 'Analysis Causality',
            'outcome_of_adverse_event': 'Outcome of Adverse Event',
        }
    
    def load_patient_data(self, file_path: str) -> pd.DataFrame:
        """Load patient data from Excel file"""
        try:
            df = pd.read_excel(file_path)
            logger.info(f"Loaded {len(df)} records from {file_path}")
            return df
        except Exception as e:
            logger.error(f"Error loading patient data: {e}")
            raise
    
    def load_template(self, template_path: str) -> str:
        """Load template from Word document"""
        try:
            doc = Document(template_path)
            template_text = ""
            
            for paragraph in doc.paragraphs:
                template_text += paragraph.text + "\n"
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        template_text += cell.text + "\n"
            
            logger.info(f"Loaded template from {template_path}")
            return template_text
        except Exception as e:
            logger.error(f"Error loading template: {e}")
            raise
    
    def extract_placeholders(self, template_text: str) -> List[str]:
        """Extract placeholders from template text"""
        placeholders = re.findall(r'\{\{([^}]+)\}\}', template_text)
        logger.info(f"Found {len(placeholders)} placeholders: {placeholders}")
        return placeholders
    
    def prepare_patient_context(self, patient_row: pd.Series) -> str:
        """Convert patient data row to context string"""
        context_parts = []
        
        for column, value in patient_row.items():
            if pd.notna(value):
                context_parts.append(f"{column}: {value}")
        
        return "\n".join(context_parts)
    
    def map_data_to_placeholders(self, patient_row: pd.Series) -> Dict[str, str]:
        """Map adae.xlsx columns to template placeholders"""
        mapped_data = {}
        
        for placeholder, column_name in self.column_mapping.items():
            if column_name in patient_row:
                value = patient_row[column_name]
                if pd.notna(value):
                    mapped_data[placeholder] = str(value)
                else:
                    mapped_data[placeholder] = "Not specified"
            else:
                mapped_data[placeholder] = "Not available"
        
        return mapped_data
    
    def generate_ai_content(self, placeholder: str, patient_context: str, mapped_data: Dict[str, str]) -> str:
        """Use AI to generate content for placeholders that need elaboration"""
        
        # Some placeholders need AI generation, others can use direct data
        elaboration_needed = [
            'clinical_course_description',
            'discussion'
        ]
        
        if placeholder not in elaboration_needed:
            return mapped_data.get(placeholder, "Not available")
        
        try:
            prompt = f"""
You are a medical writer creating an adverse event narrative. Based on the patient data provided, generate professional content for: {placeholder}

Patient Data:
{patient_context}

Mapped Data:
{json.dumps(mapped_data, indent=2)}

Instructions:
- Generate professional, accurate medical content
- Use the patient data provided
- Maintain consistency with medical writing standards
- Keep content concise but comprehensive
- For clinical course: describe the progression and management of the event
- For discussion: analyze the relationship to treatment and clinical significance

Generate content for: {placeholder}
"""
            
            response = self.client.chat.completions.create(
                model="gpt-3.5-turbo",
                messages=[
                    {"role": "system", "content": "You are a professional medical writer specializing in adverse event narratives and regulatory documentation."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=500,
                temperature=0.3
            )
            
            generated_content = response.choices[0].message.content.strip()
            logger.info(f"Generated content for {placeholder}")
            return generated_content
            
        except Exception as e:
            logger.error(f"Error generating AI content for {placeholder}: {e}")
            return mapped_data.get(placeholder, "Not available")
    
    def populate_template(self, template_text: str, patient_data: pd.DataFrame, patient_index: int = 0) -> str:
        """Populate template with patient data"""
        if patient_index >= len(patient_data):
            raise ValueError(f"Patient index {patient_index} out of range")
        
        patient_row = patient_data.iloc[patient_index]
        placeholders = self.extract_placeholders(template_text)
        
        # Map data to placeholders
        mapped_data = self.map_data_to_placeholders(patient_row)
        
        # Generate AI content for placeholders that need elaboration
        patient_context = self.prepare_patient_context(patient_row)
        
        populated_template = template_text
        
        # Replace placeholders
        for placeholder in placeholders:
            if placeholder in mapped_data:
                content = mapped_data[placeholder]
            else:
                content = self.generate_ai_content(placeholder, patient_context, mapped_data)
            
            populated_template = populated_template.replace(f"{{{{{placeholder}}}}}", content)
        
        return populated_template
    
    def save_populated_template(self, populated_text: str, output_path: str):
        """Save populated template to Word document"""
        try:
            doc = Document()
            
            paragraphs = populated_text.split('\n')
            for para_text in paragraphs:
                if para_text.strip():
                    doc.add_paragraph(para_text)
            
            doc.save(output_path)
            logger.info(f"Saved populated template to {output_path}")
        except Exception as e:
            logger.error(f"Error saving populated template: {e}")
            raise
    
    def process_patient(self, data_file: str, template_file: str, output_file: str, patient_index: int = 0):
        """Complete workflow"""
        logger.info(f"Processing patient {patient_index} from {data_file}")
        
        patient_data = self.load_patient_data(data_file)
        template_text = self.load_template(template_file)
        
        populated_text = self.populate_template(template_text, patient_data, patient_index)
        
        self.save_populated_template(populated_text, output_file)
        
        logger.info(f"Successfully processed patient {patient_index}. Output saved to {output_file}")


def main():
    """Example usage"""
    api_key = os.getenv('OPENAI_API_KEY')
    
    try:
        populator = EnhancedAITemplatePopulator(api_key=api_key)
        
        data_file = "/Users/anyasikri/Desktop/reference_docs/adae.xlsx"
        template_file = "/Users/anyasikri/Desktop/reference_docs/improved_template.docx"
        output_file = "/Users/anyasikri/Desktop/reference_docs/enhanced_narrative.docx"
        
        populator.process_patient(
            data_file=data_file,
            template_file=template_file,
            output_file=output_file,
            patient_index=0
        )
        
        print("\n✓ Success! Enhanced narrative created!")
        
    except Exception as e:
        logger.error(f"Error in main execution: {e}")
        print(f"Error: {e}")


if __name__ == "__main__":
    main()


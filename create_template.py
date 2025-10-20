#!/usr/bin/env python3
"""
Create a sample Word template for testing the AI Template Populator
"""

from docx import Document
from docx.shared import Inches

def create_sample_template():
    """
    Create a sample Word template with placeholders for patient narrative
    """
    doc = Document()
    
    # Title
    title = doc.add_heading('Patient Safety Narrative Report', 0)
    
    # Patient Information Section
    doc.add_heading('Patient Information', level=1)
    doc.add_paragraph('Patient ID: {{patient_id}}')
    doc.add_paragraph('Age: {{age}}')
    doc.add_paragraph('Gender: {{gender}}')
    
    # Medical History Section
    doc.add_heading('Medical History', level=1)
    doc.add_paragraph('{{medical_history_summary}}')
    
    # Current Medications
    doc.add_heading('Current Medications', level=1)
    doc.add_paragraph('{{current_medications}}')
    
    # Adverse Event Details
    doc.add_heading('Adverse Event Details', level=1)
    doc.add_paragraph('Event: {{adverse_event}}')
    doc.add_paragraph('Date: {{event_date}}')
    doc.add_paragraph('Severity: {{severity}}')
    doc.add_paragraph('{{event_description}}')
    
    # Clinical Assessment
    doc.add_heading('Clinical Assessment', level=1)
    doc.add_paragraph('{{clinical_assessment}}')
    
    # Outcome
    doc.add_heading('Outcome', level=1)
    doc.add_paragraph('{{outcome_description}}')
    
    # Safety Analysis
    doc.add_heading('Safety Analysis', level=1)
    doc.add_paragraph('{{safety_analysis}}')
    
    # Recommendations
    doc.add_heading('Recommendations', level=1)
    doc.add_paragraph('{{recommendations}}')
    
    # Save the template
    doc.save('sample_template.docx')
    print("Sample template created: sample_template.docx")

if __name__ == "__main__":
    create_sample_template()

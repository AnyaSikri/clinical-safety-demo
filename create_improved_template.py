#!/usr/bin/env python3
"""
Create an improved Word template for AE narratives using SDTM data
"""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_improved_template():
    """
    Create an improved template leveraging SDTM data structure
    """
    doc = Document()
    
    # Title
    title = doc.add_heading('Adverse Event Narrative Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add study header
    doc.add_paragraph('').alignment = WD_ALIGN_PARAGRAPH.CENTER
    study_info = doc.add_paragraph('Study: {{study_identifier}}')
    study_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('')
    
    # PATIENT INFORMATION SECTION
    doc.add_heading('1. Patient Information', level=1)
    
    # Demographics table
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Light Grid Accent 1'
    
    table.cell(0, 0).text = 'Subject Identifier'
    table.cell(0, 1).text = '{{unique_subject_identifier}}'
    
    table.cell(1, 0).text = 'Age'
    table.cell(1, 1).text = '{{age}} {{age_units}}'
    
    table.cell(2, 0).text = 'Sex'
    table.cell(2, 1).text = '{{sex}}'
    
    table.cell(3, 0).text = 'Race'
    table.cell(3, 1).text = '{{race}}'
    
    table.cell(4, 0).text = 'Ethnicity'
    table.cell(4, 1).text = '{{ethnicity}}'
    
    table.cell(5, 0).text = 'Study Site'
    table.cell(5, 1).text = '{{study_site_identifier}}'
    
    doc.add_paragraph('')
    
    # TREATMENT INFORMATION
    doc.add_heading('2. Treatment Information', level=1)
    
    tx_table = doc.add_table(rows=4, cols=2)
    tx_table.style = 'Light Grid Accent 1'
    
    tx_table.cell(0, 0).text = 'Study Phase'
    tx_table.cell(0, 1).text = '{{phase}}'
    
    tx_table.cell(1, 0).text = 'Planned Treatment'
    tx_table.cell(1, 1).text = '{{description_of_planned_arm}}'
    
    tx_table.cell(2, 0).text = 'Actual Treatment'
    tx_table.cell(2, 1).text = '{{description_of_actual_arm}}'
    
    tx_table.cell(3, 0).text = 'Treatment Dates'
    tx_table.cell(3, 1).text = 'First Exposure: {{datetime_of_first_exposure_to_treatment}}\nLast Exposure: {{datetime_of_last_exposure_to_treatment}}'
    
    doc.add_paragraph('')
    
    # ADVERSE EVENT DETAILS
    doc.add_heading('3. Adverse Event Details', level=1)
    
    ae_table = doc.add_table(rows=8, cols=2)
    ae_table.style = 'Light Grid Accent 1'
    
    ae_table.cell(0, 0).text = 'Reported Term'
    ae_table.cell(0, 1).text = '{{reported_term_for_the_adverse_event}}'
    
    ae_table.cell(1, 0).text = 'MedDRA Preferred Term'
    ae_table.cell(1, 1).text = '{{dictionary_derived_term}} ({{preferred_term_code}})'
    
    ae_table.cell(2, 0).text = 'Body System'
    ae_table.cell(2, 1).text = '{{body_system_or_organ_class}}'
    
    ae_table.cell(3, 0).text = 'Start Date'
    ae_table.cell(3, 1).text = '{{start_datetime_of_adverse_event}} (Study Day {{analysis_start_relative_day}})'
    
    ae_table.cell(4, 0).text = 'End Date'
    ae_table.cell(4, 1).text = '{{end_datetime_of_adverse_event}} (Study Day {{analysis_end_relative_day}})'
    
    ae_table.cell(5, 0).text = 'Serious Event'
    ae_table.cell(5, 1).text = '{{serious_event}}'
    
    ae_table.cell(6, 0).text = 'Severity Grade'
    ae_table.cell(6, 1).text = '{{standard_toxicity_grade}}'
    
    ae_table.cell(7, 0).text = 'Treatment Emergent'
    ae_table.cell(7, 1).text = '{{treatment_emergent_flag}}'
    
    doc.add_paragraph('')
    
    # SERIOUSNESS CRITERIA
    doc.add_heading('4. Seriousness Criteria', level=1)
    
    serious_table = doc.add_table(rows=5, cols=1)
    serious_table.style = 'Light Grid Accent 1'
    
    serious_table.cell(0, 0).text = 'Death: {{results_in_death}}'
    serious_table.cell(1, 0).text = 'Life-threatening: {{is_life_threatening}}'
    serious_table.cell(2, 0).text = 'Requires/Prolongs Hospitalization: {{requires_or_prolongs_hospitalization}}'
    serious_table.cell(3, 0).text = 'Persistent Disability: {{persist_or_signif_disabilityincapacity}}'
    serious_table.cell(4, 0).text = 'Other Medically Important: {{other_medically_important_serious_event}}'
    
    doc.add_paragraph('')
    
    # ACTION TAKEN AND CAUSALITY
    doc.add_heading('5. Action Taken and Causality', level=1)
    
    action_table = doc.add_table(rows=3, cols=2)
    action_table.style = 'Light Grid Accent 1'
    
    action_table.cell(0, 0).text = 'Action Taken'
    action_table.cell(0, 1).text = '{{action_taken_with_study_treatment}}'
    
    action_table.cell(1, 0).text = 'Causality Assessment'
    action_table.cell(1, 1).text = '{{causality}}'
    
    action_table.cell(2, 0).text = 'Analysis Causality'
    action_table.cell(2, 1).text = '{{analysis_causality}}'
    
    doc.add_paragraph('')
    
    # OUTCOME
    doc.add_heading('6. Outcome', level=1)
    
    doc.add_paragraph('{{outcome_of_adverse_event}}')
    
    doc.add_paragraph('')
    
    # CLINICAL COURSE AND MANAGEMENT
    doc.add_heading('7. Clinical Course and Management', level=1)
    doc.add_paragraph('{{clinical_course_description}}')
    
    doc.add_paragraph('')
    
    # DISCUSSION
    doc.add_heading('8. Discussion', level=1)
    doc.add_paragraph('{{discussion}}')
    
    doc.add_paragraph('')
    
    # Save the template
    doc.save('/Users/anyasikri/Desktop/reference_docs/improved_template.docx')
    print("Improved template created: improved_template.docx")

if __name__ == "__main__":
    create_improved_template()

